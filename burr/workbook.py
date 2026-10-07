"""Compile DSL ASTs to Excel, independently recalculate, verify, and cache results.

The calculator consumes the serialized workbook's Excel formulas, not the DSL.
It implements exactly the arithmetic, references, and functions emitted here.
"""

import ast
import datetime
import io
import math
import operator
import re
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.properties import CalcProperties

from .errors import BurrError, require
from .storage import atomic_write
from .values import number, series

BLUE, BLACK, GREEN = "0000FF", "000000", "008000"


class ExcelCalculator:
    """A separate cell-graph evaluator; no calls to the model evaluator."""

    def __init__(self, workbook):
        self.workbook, self.cache, self.active = workbook, {}, set()

    def cell(self, sheet, address):
        address = address.replace("$", "")
        key = sheet, address
        if key in self.cache:
            return self.cache[key]
        require(key not in self.active, "eval_error", f"spreadsheet circular reference: {sheet}!{address}")
        self.active.add(key)
        try:
            value = self.workbook[sheet][address].value
            result = self.formula(value, sheet) if isinstance(value, str) and value.startswith("=") else number(value)
            self.cache[key] = number(result)
            return self.cache[key]
        finally:
            self.active.remove(key)

    def reference(self, token, sheet):
        if "!" in token:
            sheet, token = token.rsplit("!", 1)
            sheet = sheet.strip("'").replace("''", "'")
        if ":" not in token:
            require(re.fullmatch(r"\$?[A-Z]+\$?\d+", token), "eval_error", f"unsupported spreadsheet reference {token}")
            return self.cell(sheet, token)
        left, top, right, bottom = range_boundaries(token.replace("$", ""))
        return [self.cell(sheet, f"{get_column_letter(c)}{r}") for r in range(top, bottom + 1) for c in range(left, right + 1)]

    def formula(self, formula, sheet):
        tokens = [t for t in Tokenizer(formula).items if t.type != "WHITE-SPACE"]
        index = 0
        operations = {"+": (10, operator.add), "-": (10, operator.sub), "*": (20, operator.mul),
                      "/": (20, operator.truediv), "^": (30, operator.pow)}

        def expression(minimum=0):
            nonlocal index
            require(index < len(tokens), "eval_error", "incomplete Excel formula")
            token = tokens[index]
            index += 1
            if token.type == "OPERAND":
                left = float(token.value) if token.subtype == "NUMBER" else self.reference(token.value, sheet)
            elif token.type == "OPERATOR-PREFIX" and token.value in {"-", "+"}:
                left = expression(40)
                if token.value == "-":
                    left = -left
            elif token.type == "PAREN" and token.subtype == "OPEN":
                left = expression()
                require(index < len(tokens) and tokens[index].type == "PAREN" and tokens[index].subtype == "CLOSE", "eval_error", "unclosed Excel parentheses")
                index += 1
            elif token.type == "FUNC" and token.subtype == "OPEN":
                name, arguments = token.value[:-1].upper(), []
                while True:
                    arguments.append(expression())
                    require(index < len(tokens), "eval_error", "unclosed Excel function")
                    if tokens[index].type == "SEP":
                        index += 1
                        continue
                    require(tokens[index].type == "FUNC" and tokens[index].subtype == "CLOSE", "eval_error", "invalid Excel function arguments")
                    index += 1
                    break
                flattened = [x for arg in arguments for x in (arg if isinstance(arg, list) else [arg])]
                require(name in {"SUM", "MIN", "MAX", "ABS"}, "eval_error", f"unsupported Excel function {name}")
                left = {"SUM": sum, "MIN": min, "MAX": max, "ABS": lambda xs: abs(xs[0])}[name](flattened)
            else:
                raise BurrError("eval_error", f"unsupported Excel token {token.value!r}")
            while index < len(tokens) and tokens[index].type == "OPERATOR-INFIX":
                op = tokens[index].value
                require(op in operations, "eval_error", f"unsupported Excel operator {op}")
                precedence, function = operations[op]
                if precedence < minimum:
                    break
                index += 1
                right = expression(precedence + (0 if op == "^" else 1))
                left = function(left, right)
            return left

        result = expression()
        require(index == len(tokens), "eval_error", "unparsed Excel tokens")
        return number(result)

    def recalculate(self):
        try:
            for sheet in self.workbook:
                for row in sheet:
                    for cell in row:
                        if cell.data_type == "f":
                            self.cell(sheet.title, cell.coordinate)
        except (ArithmeticError, ValueError, KeyError, RecursionError) as exc:
            raise BurrError("eval_error", f"spreadsheet recalculation failed: {exc}") from exc
        return self.cache


class Compiler:
    def __init__(self, model):
        self.model = model
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)
        for name in ("Assumptions", "Model", "DCF", "Actuals"):
            self.workbook.create_sheet(name)
        self.input_rows, self.line_rows, self.grow_rows = {}, {}, {}

    def formula(self, node, t):
        col = get_column_letter(t + 2)
        if isinstance(node, ast.Constant):
            return str(node.value)
        if isinstance(node, ast.Name):
            if node.id in self.input_rows:
                return f"Assumptions!{col}{self.input_rows[node.id]}"
            return f"{col}{self.line_rows[node.id]}"
        if isinstance(node, ast.UnaryOp):
            return f"(-{self.formula(node.operand, t)})"
        if isinstance(node, ast.BinOp):
            op = {ast.Add: "+", ast.Sub: "-", ast.Mult: "*", ast.Div: "/"}[type(node.op)]
            return f"({self.formula(node.left, t)}{op}{self.formula(node.right, t)})"
        name, args = node.func.id, node.args
        if name == "grow":
            return f"({self.formula(args[0], 0)}*{col}{self.grow_rows[node]})"
        if name in {"lag", "delta"}:
            previous = self.formula(args[0], t - 1) if t else self.formula(args[1], 0)
            return previous if name == "lag" else f"({self.formula(args[0], t)}-{previous})"
        values = [self.formula(arg, t) for arg in args]
        if name == "clip":
            return f"MIN(MAX({values[0]},{values[1]}),{values[2]})"
        fn = {"minimum": "MIN", "maximum": "MAX", "abs": "ABS"}[name]
        return f"{fn}({','.join(values)})"

    def build(self):
        m, wb = self.model, self.workbook
        m.value()
        assumptions, model, dcf, actuals = [wb[n] for n in ("Assumptions", "Model", "DCF", "Actuals")]
        for sheet in (assumptions, model, dcf):
            sheet.append(["Line", *m.years])
        n = len(m.years)
        for name, value in sorted(m.inputs.items()):
            self.input_rows[name] = assumptions.max_row + 1
            assumptions.append([name, *series(value, n)])
            form = m.params["bindings"].get(name, m.workspace.world.get(name))
            if isinstance(form, dict):
                from openpyxl.comments import Comment
                metadata = "\n".join(f"{key}: {form[key]}" for key in ("rationale", "source", "confidence", "as_of", "unit") if key in form)
                if metadata:
                    assumptions.cell(assumptions.max_row, 1).comment = Comment(metadata, "burr")
        valuation_rows = {}
        for name in ("wacc", "terminal_growth", "net_cash", "shares"):
            valuation_rows[name] = assumptions.max_row + 1
            assumptions.append([f"valuation.{name}", *series(m.valuation.get(name, 0.0), n)])
        for name in sorted(m.expressions):
            self.line_rows[name] = model.max_row + 1
            model.append([name])
        for name, target in m.valuation_refs.items():
            sheet, row = (("Assumptions", self.input_rows[target]) if target in self.input_rows
                          else ("Model", self.line_rows[target]))
            for t in range(n):
                assumptions.cell(valuation_rows[name], t + 2, f"={sheet}!{get_column_letter(t + 2)}{row}")
        helper = model.max_row + 1
        for name in sorted(m.expressions):
            for node in ast.walk(m.expressions[name]):
                if isinstance(node, ast.Call) and node.func.id == "grow":
                    self.grow_rows[node] = helper
                    model.cell(helper, 1, f"_grow_factor_{len(self.grow_rows)}")
                    model.row_dimensions[helper].hidden = True
                    helper += 1
        for name, row in self.line_rows.items():
            for t in range(n):
                model.cell(row, t + 2, "=" + self.formula(m.expressions[name], t))
        for node, row in self.grow_rows.items():
            for t in range(n):
                previous = f"{get_column_letter(t + 1)}{row}" if t else "1"
                model.cell(row, t + 2, f"=({previous}*(1+{self.formula(node.args[1], t)}))")
        for row, label in enumerate(("FCF", "WACC", "Discount factor", "Present value"), 2):
            dcf.cell(row, 1, label)
        for t in range(n):
            col, previous = get_column_letter(t + 2), get_column_letter(t + 1)
            dcf.cell(2, t + 2, f"=Model!{col}{self.line_rows[m.params['valuation']['fcf']]}")
            dcf.cell(3, t + 2, f"=Assumptions!{col}{valuation_rows['wacc']}")
            dcf.cell(4, t + 2, f"={(previous + '4') if t else '1'}*(1+{col}3)")
            dcf.cell(5, t + 2, f"={col}2/{col}4")
        last = get_column_letter(n + 1)
        g = f"Assumptions!{last}{valuation_rows['terminal_growth']}"
        summaries = [
            (7, "Terminal value", f"={last}2*(1+{g})/({last}3-{g})"),
            (8, "Terminal present value", f"=B7/{last}4"),
            (9, "Enterprise value", f"=SUM(B5:{last}5)+B8"),
            (10, "Net cash", f"=Assumptions!B{valuation_rows['net_cash']}"),
            (11, "Equity value", "=B9+B10"),
            (12, "Shares", f"=Assumptions!B{valuation_rows['shares']}"),
            (13, "Per-share value", "=B11/B12"),
            (14, "Terminal weight", "=B8/B9" if m.value()["ev"] else "=0"),
        ]
        for row, label, formula in summaries:
            dcf.cell(row, 1, label)
            dcf.cell(row, 2, formula)
        actuals.append(["period", "line", "value"])
        for row in m.workspace.actuals:
            actuals.append([row["period"], row["line"], row["value"]])
        actuals.protection.sheet = True
        actuals.protection.set_password("burr-actuals")
        for sheet in wb:
            sheet.freeze_panes = "B2"
            sheet.column_dimensions["A"].width = 30
            sheet.sheet_view.showGridLines = False
            for col in range(2, max(n + 2, 4)):
                sheet.column_dimensions[get_column_letter(col)].width = 18
            for row in sheet:
                for cell in row:
                    cell.font = Font(name="Aptos", size=11, color=BLACK)
                    if cell.row == 1:
                        cell.fill = PatternFill("solid", fgColor="16324F")
                        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
                    elif cell.data_type == "f":
                        cell.font = Font(name="Aptos", color=GREEN if "!" in cell.value else BLACK)
                        cell.number_format = "#,##0.00;[Red](#,##0.00)"
                    elif sheet.title == "Assumptions" and cell.column > 1:
                        cell.font = Font(name="Aptos", color=BLUE)
                        cell.protection = Protection(locked=False)
                        cell.number_format = "0.0000"
            sheet.auto_filter.ref = sheet.dimensions if sheet.title == "Actuals" else "A1:A1"
        wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
        wb.properties.creator = "burr"
        wb.properties.title = f"{m.entity} — {m.scenario}"
        wb.properties.created = wb.properties.modified = datetime.datetime(2000, 1, 1)
        return wb


def verify(model, workbook, line_rows):
    calculator = ExcelCalculator(workbook)
    cache = calculator.recalculate()
    for name, values in model.run().items():
        for t, expected in enumerate(values):
            address = f"{get_column_letter(t + 2)}{line_rows[name]}"
            actual = calculator.cell("Model", address)
            require(math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-9), "eval_error",
                    f"spreadsheet disagrees at Model!{address}: workbook={actual}, engine={expected}")
    for key, row in {"ev": 9, "equity": 11, "per_share": 13, "terminal_weight": 14}.items():
        actual, expected = calculator.cell("DCF", f"B{row}"), model.value()[key]
        require(math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-9), "eval_error",
                f"spreadsheet DCF {key} disagrees: workbook={actual}, engine={expected}")
    return cache


def cached_archive(data, cache, sheet_names):
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", name)
            if match:
                sheet = sheet_names[int(match[1]) - 1]
                root = ET.fromstring(content)
                for cell in root.iter(f"{{{namespace}}}c"):
                    if cell.find(f"{{{namespace}}}f") is not None:
                        value = cell.find(f"{{{namespace}}}v")
                        if value is None:
                            value = ET.SubElement(cell, f"{{{namespace}}}v")
                        value.text = repr(cache[sheet, cell.attrib["r"]])
                content = ET.tostring(root, encoding="utf-8")
            elif name == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                                 rb"\g<1>2000-01-01T00:00:00Z\g<2>", content)
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, content)
    return output.getvalue()


def emit(model, path):
    compiler = Compiler(model)
    workbook = compiler.build()
    stream = io.BytesIO()
    workbook.save(stream)
    serialized = load_workbook(io.BytesIO(stream.getvalue()))
    cache = verify(model, serialized, compiler.line_rows)
    data = cached_archive(stream.getvalue(), cache, workbook.sheetnames)
    # Confirm cached numbers are actually readable from the finished artifact.
    cached = load_workbook(io.BytesIO(data), data_only=True)
    require(math.isclose(cached["DCF"]["B13"].value, model.value()["per_share"], rel_tol=1e-10),
            "eval_error", "spreadsheet cached result disagrees")
    atomic_write(path, data)
    return {"ok": True, "output": str(path), "verified": True, "sheets": workbook.sheetnames,
            "per_share": model.value()["per_share"]}
