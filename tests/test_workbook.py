import io

import pytest
from openpyxl import load_workbook

from burr.errors import BurrError
from burr.workbook import Compiler, ExcelCalculator, emit, verify


@pytest.mark.parametrize("fixture", ["workspace", "recurrence"])
def test_live_workbook_verified_and_deterministic(request, fixture, tmp_path):
    workspace = request.getfixturevalue(fixture)
    model = workspace.model()
    path = tmp_path / "model.xlsx"
    assert emit(model, path)["verified"]
    first = path.read_bytes()
    emit(workspace.model(), path)
    assert path.read_bytes() == first
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Assumptions", "Model", "DCF", "Actuals"]
    assert workbook["Actuals"].protection.sheet
    assert all(c.data_type == "f" for row in workbook["Model"].iter_rows(min_row=2, min_col=2) for c in row)
    cached = load_workbook(path, data_only=True)
    assert cached["DCF"]["B13"].value == pytest.approx(model.value()["per_share"])
    calculator = ExcelCalculator(workbook)
    original = calculator.cell("DCF", "B13")
    for row in workbook["Assumptions"]:
        if row[0].value == "price":
            for cell in row[1:]:
                cell.value *= 1.1
    changed = ExcelCalculator(workbook).cell("DCF", "B13")
    assert changed == pytest.approx(original * 1.1)


def test_disagreement_blocks_export(workspace, tmp_path, monkeypatch):
    original = Compiler.formula

    def wrong_formula(self, node, t):
        return "(" + original(self, node, t) + "+1)"

    monkeypatch.setattr(Compiler, "formula", wrong_formula)
    path = tmp_path / "model.xlsx"
    path.write_bytes(b"existing output")
    with pytest.raises(BurrError, match="disagrees"):
        emit(workspace.model(), path)
    assert path.read_bytes() == b"existing output"
