"""Offline regressions against source-workbook calculations, not Burr exports."""

import copy
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from burr.errors import BurrError
from burr.model import Workspace
from burr.storage import read_yaml, yaml_text
from burr.workbook import emit
from benchmarks.damodaran.catalog import MODELS, ROOT, variants
from benchmarks.damodaran.evaluate import compare, evaluate
from benchmarks.damodaran.exercise import check_recalculated
from benchmarks.damodaran.models import write_workspace
from benchmarks.damodaran.oracle import extract_inputs, extract_outputs, fetch, load_json


@pytest.fixture(scope="module")
def corpus(tmp_path_factory):
    root = tmp_path_factory.mktemp("damodaran") / "workspace"
    inputs = load_json(ROOT / "inputs.json")
    reference = load_json(ROOT / "reference.json")
    sources = load_json(ROOT / "sources.json")
    write_workspace(root, inputs)
    return root, inputs, reference, sources


@pytest.fixture(scope="module")
def report(corpus):
    return evaluate(*corpus)


@pytest.mark.parametrize("case", [f"{name}--{s}" for name in MODELS for s in variants(name)])
def test_original_workbook_values_and_perturbations(report, case):
    result = report["cases"][case]
    failures = [c for c in result["checks"] + result.get("delta_checks", []) if not c["ok"]]
    assert not failures, failures
    assert result["ok"]


@pytest.mark.parametrize("name", MODELS)
def test_exported_workbook_matches_external_reference(corpus, tmp_path, name):
    root, _, reference, _ = corpus
    path = tmp_path / f"{name}.xlsx"
    emit(Workspace(root / "companies" / name).model(), path)
    wb = load_workbook(path, data_only=True)
    # Stable model uses one synthetic share and no capital bridge: price = EV.
    expected = reference["cases"][f"{name}--base"]["outputs"]
    key = "ev" if name == "stable" else "per_share"
    assert wb["DCF"]["B13"].value == pytest.approx(expected[key][0], rel=1e-9, abs=1e-7)


def test_mutated_formula_is_detected(corpus, tmp_path):
    _, inputs, reference, sources = corpus
    root = write_workspace(tmp_path / "mutant", inputs)
    path = root / "templates/stable.yaml"
    template = read_yaml(path)
    template["lines"]["fcff_current"] = "nopat_current + net_capex_current - working_capital_change_current"
    path.write_text(yaml_text(template))
    result = evaluate(root, inputs, reference, sources)
    assert not result["ok"] and result["failed_checks"] > 0


@pytest.mark.parametrize("fault", ["case", "metric", "vector", "input", "source"])
def test_incomplete_or_stale_reference_is_rejected(corpus, fault):
    root, inputs, original, sources = corpus
    reference = copy.deepcopy(original)
    case = reference["cases"]["two_stage--base"]
    if fault == "case":
        del reference["cases"]["two_stage--base"]
    elif fault == "metric":
        del case["outputs"]["fcff"]
    elif fault == "vector":
        case["outputs"]["fcff"].pop()
    elif fault == "input":
        case["input_digest"] = "wrong"
    else:
        case["source_sha256"] = "wrong"
    with pytest.raises(ValueError):
        evaluate(root, inputs, reference, sources)


def test_oracle_rejects_uncalculated_cells():
    wb = Workbook()
    wb.active.title = MODELS["stable"]["sheet"]
    wb.active["F51"] = "=1+2"
    with pytest.raises(ValueError, match="oracle cell"):
        extract_outputs(wb, "stable")


def test_input_adapter_rejects_unreviewed_formulas(corpus):
    _, inputs, _, _ = corpus
    wb = Workbook()
    spec = MODELS["stable"]
    wb.active.title = spec["sheet"]
    for field, cell in spec["inputs"].items():
        wb.active[cell] = inputs["stable"][field]
    for cell, value in spec["guards"].items():
        wb.active[cell] = value
    assert extract_inputs(wb, "stable") == inputs["stable"]
    wb.active["D21"] = "=F61"
    with pytest.raises(ValueError, match="literal finite input"):
        extract_inputs(wb, "stable")


def test_corrupt_cached_source_is_rejected_without_network(tmp_path):
    path = tmp_path / "originals/fcffst.xls"
    path.parent.mkdir()
    path.write_bytes(b"not the pinned workbook")
    with pytest.raises(ValueError, match="checksum mismatch"):
        fetch(tmp_path)


def test_nonfinite_results_cannot_pass():
    with pytest.raises(ValueError, match="Non-finite"):
        compare({"ev": [float("nan")]}, {"ev": [1]}, "invalid")


def test_workspace_generation_preserves_existing_journal(corpus):
    root, inputs, _, _ = corpus
    with pytest.raises(FileExistsError):
        write_workspace(root, inputs)


def test_invalid_terminal_assumption_is_diagnostic(corpus):
    root, _, _, _ = corpus
    ws = Workspace(root / "companies/stable")
    params = copy.deepcopy(ws.params)
    params["valuation"]["terminal_growth"] = .99
    with pytest.raises(BurrError, match="terminal WACC"):
        ws.model(params=params).value()


def test_saved_examples_match_source_workbooks(corpus):
    _, inputs, reference, sources = corpus
    root = Path(__file__).resolve().parents[1] / "examples/damodaran"
    report = evaluate(root, inputs, reference, sources)
    assert report["ok"], report["failed_checks"]


def test_external_recalculation_checker_detects_changed_and_missing_values(corpus, tmp_path):
    root, _, _, _ = corpus
    original = tmp_path / "original.xlsx"
    emit(Workspace(root / "companies/stable").model(), original)
    assert check_recalculated(original, original)["ok"]
    actual = load_workbook(original, data_only=True)
    actual["DCF"]["B13"] = actual["DCF"]["B13"].value + 1
    changed = tmp_path / "changed.xlsx"
    actual.save(changed)
    result = check_recalculated(original, changed)
    assert not result["ok"]
    assert result["failures"][0]["metric"] == "DCF!B13"
    actual["DCF"]["B13"] = None
    actual.save(changed)
    with pytest.raises(ValueError, match="Non-finite/non-numeric"):
        check_recalculated(original, changed)
