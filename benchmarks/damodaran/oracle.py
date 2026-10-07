"""Fetch pinned originals and recalculate their formulas with LibreOffice."""

import hashlib
import json
import math
import os
import shutil
import subprocess
import tempfile
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .catalog import MODELS, ROOT, case_inputs, variants


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, allow_nan=False).encode()).hexdigest()


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_json(path):
    return json.loads(Path(path).read_text())


def save_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, sort_keys=True, allow_nan=False) + "\n")


def fetch(cache):
    manifest = load_json(ROOT / "sources.json")
    target = Path(cache) / "originals"
    target.mkdir(parents=True, exist_ok=True)
    for name, entry in manifest["sources"].items():
        path = target / entry["filename"]
        if not path.exists():
            with urllib.request.urlopen(entry["url"], timeout=60) as response:
                content = response.read()
            if hashlib.sha256(content).hexdigest() != entry["sha256"]:
                raise ValueError(f"Source changed for {name}; review and repin explicitly")
            path.write_bytes(content)
        if file_hash(path) != entry["sha256"]:
            raise ValueError(f"Source checksum mismatch: {path}")
    return target


def libreoffice(executable=None):
    candidates = [executable, os.environ.get("BURR_LIBREOFFICE"), shutil.which("libreoffice"),
                  shutil.which("soffice"), "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(Path(candidate).resolve())
    raise RuntimeError("LibreOffice required for live references; use --soffice or BURR_LIBREOFFICE")


def convert(executable, paths, destination):
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="burr-calc-profile-") as profile:
        command = [executable, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless",
                   "--convert-to", "xlsx", "--outdir", str(destination), *map(str, paths)]
        result = subprocess.run(command, text=True, capture_output=True, timeout=180)
        if result.returncode:
            raise RuntimeError(f"LibreOffice failed: {result.stderr}\n{result.stdout}")
    outputs = [destination / (Path(p).stem + ".xlsx") for p in paths]
    for path in outputs:
        if not path.is_file():
            raise RuntimeError(f"LibreOffice did not produce {path}: {result.stdout} {result.stderr}")
    return outputs


def extract_inputs(workbook, name):
    """Read reviewed literal inputs only, with one explicitly allowed direct link."""
    spec = MODELS[name]
    sheet = workbook[spec["sheet"]]
    for cell, expected in spec["guards"].items():
        if sheet[cell].value != expected:
            raise ValueError(f"Unsupported source branch {name}!{cell}: {sheet[cell].value!r}")
    values = {}
    for field, cell in spec["inputs"].items():
        value = sheet[cell].value
        linked = spec.get("linked_inputs", {}).get(field)
        if linked:
            target = spec["inputs"][linked]
            if value != f"={target}":
                raise ValueError(f"Expected reviewed link {cell}={target}, got {value!r}")
            value = sheet[target].value
        if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value):
            raise ValueError(f"Expected literal finite input at {name}!{cell}, got {value!r}")
        values[field] = value
    return values


def extract_outputs(workbook, name):
    sheet = workbook[MODELS[name]["sheet"]]
    result = {}
    for key, cells in MODELS[name]["outputs"].items():
        result[key] = []
        for address in cells:
            value = sheet[address].value
            if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value):
                raise ValueError(f"Missing/error/non-numeric oracle cell {name}!{address}: {value!r}")
            result[key].append(float(value))
    return result


def recalculate(cache, executable=None):
    executable = libreoffice(executable)
    originals = fetch(cache)
    version = subprocess.check_output([executable, "--version"], text=True).strip()
    manifest = load_json(ROOT / "sources.json")
    with tempfile.TemporaryDirectory(prefix="damodaran-live-", dir=Path(cache)) as tmp:
        tmp = Path(tmp)
        converted = convert(executable, [originals / m["filename"] for m in MODELS.values()], tmp / "converted")
        inputs, pending, prepared = {}, [], tmp / "prepared"
        prepared.mkdir()
        for (name, spec), original in zip(MODELS.items(), converted):
            inputs[name] = extract_inputs(load_workbook(original), name)
            for scenario, edits in variants(name).items():
                wb = load_workbook(original)
                sheet = wb[spec["sheet"]]
                for field, value in edits.items():
                    sheet[spec["inputs"][field]] = value
                path = prepared / f"{name}--{scenario}.xlsx"
                wb.save(path)
                pending.append(path)
        computed = convert(executable, pending, tmp / "recalculated")
        cases = {}
        for path, prepared_path in zip(computed, pending):
            name, scenario = path.stem.split("--")
            cases[path.stem] = {
                "model": name, "scenario": scenario,
                "input_digest": digest(case_inputs(name, inputs[name], scenario)),
                "source_sha256": manifest["sources"][name]["sha256"],
                "prepared_workbook_sha256": file_hash(prepared_path),
                "recalculated_workbook_sha256": file_hash(path),
                "outputs": extract_outputs(load_workbook(path, data_only=True), name),
            }
        return inputs, {
            "schema_version": 1, "oracle": version,
            "captured_at": datetime.now(timezone.utc).isoformat(),
            "catalog_digest": digest(MODELS), "cases": cases,
        }
