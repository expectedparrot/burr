"""Stress the saved pilot artifacts without modifying either submission.

Run: .venv/bin/python -m benchmarks.damodaran.workflow_eval --output DIR
This is a feature exercise, not a new randomized agent comparison.
"""

import argparse
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys

import yaml

from burr.model import Workspace

ROOT = Path(__file__).resolve().parent
PILOT = ROOT / "agent_pilot"


def save(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=2, sort_keys=True) + "\n")


def snapshot(root):
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in root.rglob("*") if p.is_file() and "__pycache__" not in p.parts}


def run(destination, linked_valuation=False):
    destination = destination.resolve()
    destination.mkdir(parents=True, exist_ok=False)
    plan = {
        "kind": "deterministic stress tests of existing pilot artifacts; no new agents",
        "criteria": [
            "Read-only exploration preserves files and matches direct Python values.",
            "A named scenario preserves the base and records source/rationale.",
            "A source-only revision invalidates evidence even when numbers hold.",
            "A numerical revision marks drift; recovery appends rather than overwrites evidence.",
            "An edited Excel export is detected and regenerated with a current receipt.",
            "A risk-free binding edit exposes any stale dependent valuation assumptions.",
        ],
        "python_comparator": "Unmodified prior submission; no added workflow safeguards.",
        "interpretation": "Absence of a Python safeguard is not inability to implement one.",
        "valuation_migration": "explicit references" if linked_valuation else "original copied values",
    }
    save(destination / "plan.json", plan)
    original = snapshot(PILOT / "submissions")
    spec = importlib.util.spec_from_file_location("pilot_python", PILOT / "submissions/python/submission.py")
    python = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(python)
    baseline = json.loads((PILOT / "baseline.json").read_text())
    expected_base = python.model(baseline)["per_share"][0]
    heldout_inputs = json.loads((PILOT / "hidden-inputs.json").read_text())
    oracle = json.loads((PILOT / "oracle.json").read_text())

    def source_value(inputs):
        case = next(k for k, values in heldout_inputs.items() if values == inputs)
        return oracle["cases"][case]["per_share"][0]

    if not math.isclose(expected_base, source_value(baseline), rel_tol=1e-9, abs_tol=1e-7):
        raise ValueError("Python baseline disagrees with saved LibreOffice oracle")
    transcript = []

    def cli(argv, cwd=None):
        p = subprocess.run(list(map(str, argv)), cwd=cwd, capture_output=True, text=True, timeout=60)
        if p.returncode:
            raise RuntimeError(p.stdout + p.stderr)
        obj = json.loads(p.stdout)
        transcript.append({"argv": list(map(str, argv)), "result": obj})
        return obj

    def fresh(name):
        root = destination / name
        if linked_valuation:
            from .linked_pilot import create
            create(root)
        else:
            shutil.copytree(PILOT / "submissions/burr/model", root)
        inspect(root)
        return root

    def inspect(root, *args):
        return cli([sys.executable, "-m", "burr", "--json", "-C", root,
                    "next", "fcff", "--deliverable", "excel", *args])

    def action(step):
        cli(step["action"]["argv"], step["action"]["cwd"])
        return cli(step["resume_argv"])

    def finish(step):
        for _ in range(4):
            if step["complete"] or step["action"] is None:
                return step
            step = action(step)
        raise RuntimeError("Unexpected guidance loop")

    def edit(root, key, field, value):
        path = root / "companies/fcff/params.yaml"
        data = yaml.safe_load(path.read_text())
        data["bindings"][key][field] = value
        path.write_text(yaml.safe_dump(data, sort_keys=False))
        return inspect(root)

    def model(root, scenario="base"):
        return Workspace(root / "companies/fcff").model(scenario)

    def close(a, b):
        return math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-7)

    results = {}
    root = fresh("01-explore")
    patch = root / "proposal.yaml"
    proposed = {"bindings": {"stable_roc": {"value": .147,
        "source": "Synthetic analyst hypothesis, 2026-10-07",
        "rationale": "Higher stable capital efficiency, all other inputs fixed."}}}
    patch.write_text(yaml.safe_dump(proposed))
    before = snapshot(root)
    step = inspect(root, "--patch", patch, "--question", "What if stable ROC is 14.7%?")
    trial = step["evidence"]["trial"]
    trial_inputs = {**baseline, "stable_roc": .147}
    python_inputs_before = dict(trial_inputs)
    expected_trial = python.model(trial_inputs)["per_share"][0]
    results["read_only_exploration"] = {
        "pass": snapshot(root) == before and close(trial["trial_per_share"], expected_trial)
            and close(expected_trial, source_value(trial_inputs)) and trial_inputs == python_inputs_before,
        "files_unchanged": snapshot(root) == before,
        "burr_trial": trial["trial_per_share"], "python_trial": expected_trial,
        "libreoffice_reference": source_value(trial_inputs),
        "python_input_unchanged": trial_inputs == python_inputs_before,
        "base": expected_base, "attribution": trial["attribution"],
        "python_observation": "Pure model(input_dict) also preserves its input; both support numerical exploration.",
    }
    base_before = model(root).value()["per_share"]
    params_before = yaml.safe_load((root / "companies/fcff/params.yaml").read_text())
    durable = root / "scenario.yaml"
    durable.write_text(yaml.safe_dump({"op": "add_scenario", "name": "efficient",
        "thesis": "Analyst hypothesis: improved stable capital efficiency.", **proposed}))
    cli([sys.executable, "-m", "burr", "--json", "-C", root, "patch", "fcff", "-f", durable])
    inspect(root)
    params_after = yaml.safe_load((root / "companies/fcff/params.yaml").read_text())
    results["named_scenario"] = {
        "pass": close(model(root).value()["per_share"], base_before)
            and params_before["bindings"] == params_after["bindings"]
            and close(model(root, "efficient").value()["per_share"], expected_trial),
        "base_bindings_unchanged": params_before["bindings"] == params_after["bindings"],
        "saved_scenario": params_after["scenarios"]["efficient"],
        "python_observation": "Separate input JSON files work, but naming, sources and rationale need an additional convention.",
    }

    root = fresh("02-source-revision")
    before_value = model(root).value()["per_share"]
    step = edit(root, "stable_roc", "source", "Synthetic analyst memo revised 2026-10-07; numeric estimate unchanged")
    results["source_only_revision"] = {
        "pass": not step["complete"] and step["evidence"]["replay"]["status"] == "HOLDS"
            and any(w["code"] == "stale_evidence" for w in step["warnings"])
            and close(before_value, model(root).value()["per_share"]),
        "state": step["state"], "replay": step["evidence"]["replay"]["status"],
        "warnings": step["warnings"],
        "python_observation": "The numeric input/output contract has no source metadata or evidence-freshness check.",
    }

    root = fresh("03-revision-and-handoff")
    journal = root / "companies/fcff/experiments/exp-001.yaml"
    old = journal.read_bytes()
    step = edit(root, "capex0", "value", 2673)
    drift = step["evidence"]["replay"]
    recovery_patch = root / "recovered-proposal.yaml"
    recovery_patch.write_text(yaml.safe_dump(step["evidence"]["suggested_patch"]))
    replay_expected = python.model({**baseline, "capex0": 2673, "stable_roc": .132})["per_share"][0]
    step = inspect(root, "--patch", recovery_patch, "--question", yaml.safe_load(old)["question"],
        "--finding", "Replayed the original stable-ROC hypothesis after the synthetic capex revision; value still rises.")
    completed = finish(step)
    # A handoff gets only the saved workspace, without a conversation or proposal file.
    handoff = destination / "04-handoff-only-workspace"
    shutil.copytree(root, handoff, ignore=shutil.ignore_patterns("recovered-proposal.yaml", "artifacts"))
    handed = inspect(handoff)
    results["revision_replay_and_handoff"] = {
        "pass": drift["status"] == "DRIFTED" and close(drift["current"]["trial_per_share"], replay_expected)
            and completed["complete"] and journal.read_bytes() == old
            and handed["evidence"]["replay"]["status"] == "HOLDS",
        "initial_replay": drift["status"], "recorded_trial": drift["recorded"]["trial_per_share"],
        "revised_trial": drift["current"]["trial_per_share"], "python_trial": replay_expected,
        "old_record_unchanged": journal.read_bytes() == old,
        "new_record": completed["evidence"]["selected_experiment"],
        "handoff_replay_without_proposal_file": handed["evidence"]["replay"]["status"],
        "python_observation": "Recalculation works; the submission has no persistent hypothesis, finding, or replay journal.",
    }

    root = fresh("05-edited-export")
    step = inspect(root)
    if step["state"] == "needs_input":
        # A migration invalidates old fingerprints, even when the result holds.
        proposal = root / "reviewed-proposal.yaml"
        proposal.write_text(yaml.safe_dump(step["evidence"]["suggested_patch"]))
        step = inspect(root, "--patch", proposal, "--question", "Does the saved sensitivity still hold after migration?",
                       "--finding", "Reviewed the replay: the stable-ROC sensitivity still holds after linking valuation assumptions.")
    step = finish(step)
    from openpyxl import load_workbook
    output = Path(step["evidence"]["artifact"]["output"])
    wb = load_workbook(output)
    wb["DCF"]["B13"] = 999
    wb.save(output)
    detected = inspect(root)
    repaired = finish(detected)
    results["modified_export"] = {
        "pass": detected["state"] == "export"
            and detected["evidence"]["artifact"]["reason"] == "stale_or_modified_output"
            and repaired["complete"]
            and close(load_workbook(output, data_only=True)["DCF"]["B13"].value, expected_base),
        "detected_reason": detected["evidence"]["artifact"]["reason"],
        "repaired_complete": repaired["complete"],
        "python_observation": "The saved Python workbook has no digest receipt or command to check artifact freshness.",
    }

    root = fresh("06-dependent-assumption")
    step = edit(root, "riskfree", "value", .0707)
    m = model(root)
    m.validate()
    native = m.value()["per_share"]
    explicit = m.run()["per_share"][4]
    expected = python.model({**baseline, "riskfree": .0707})["per_share"][0]
    reference = source_value({**baseline, "riskfree": .0707})
    results["dependent_assumption_consistency"] = {
        "pass": close(native, expected),
        "burr_native_dcf": native, "burr_formula_line": explicit, "python": expected,
        "libreoffice_reference": reference,
        "python_matches_source": close(expected, reference),
        "absolute_error": abs(native - expected),
        "native_validation_passed": True, "next_state": step["state"],
        "replay_status": step["evidence"]["replay"]["status"],
        "diagnostics": step["diagnostics"], "warnings": step["warnings"],
        "interpretation": ("Linked valuation updates with the underlying input; no regeneration is required."
                           if close(native, reference) else
                           "Burr flags stale evidence, but does not diagnose the adapter's duplicated WACC. Regeneration is required."),
    }
    assert snapshot(PILOT / "submissions") == original, "Archived submissions changed"
    report = {"plan": plan, "results": results, "archived_submissions_unchanged": True,
        "passed": sum(r["pass"] for r in results.values()), "total": len(results),
        "python_comparison_limit": "Observations are based on the existing code contract, not a new agent tasked with adding safeguards."}
    save(destination / "results.json", report)
    save(destination / "transcript.json", transcript)
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--linked-valuation", action="store_true", help="Evaluate a corrected copy using valuation references")
    args = parser.parse_args()
    result = run(args.output, args.linked_valuation)
    print(json.dumps(result, indent=2))
