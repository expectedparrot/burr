"""Recalculate the source on held-out inputs, then score frozen submissions.

Usage: python -m benchmarks.damodaran.agent_pilot.grade oracle RUN_DIR
       python -m benchmarks.damodaran.agent_pilot.grade score RUN_DIR
"""

import argparse
import hashlib
import json
import math
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from benchmarks.damodaran.catalog import MODELS
from benchmarks.damodaran.oracle import convert, digest, extract_outputs, libreoffice, load_json, save_json

SPEC = MODELS["two_stage"]
REL_TOL, ABS_TOL = 1e-9, 1e-7


def make_oracle(root):
    cases = load_json(root / "hidden-inputs.json")
    with tempfile.TemporaryDirectory(prefix="pilot-oracle-") as tmp:
        tmp = Path(tmp)
        pending = []
        for name, inputs in cases.items():
            wb = load_workbook(root / "burr/source.xlsx")
            sheet = wb[SPEC["sheet"]]
            for field, cell in SPEC["inputs"].items():
                sheet[cell] = inputs[field]
            path = tmp / f"{name}.xlsx"
            wb.save(path)
            pending.append(path)
        executable = libreoffice()
        recalculated = convert(executable, pending, tmp / "recalculated")
        outputs = {p.stem: extract_outputs(load_workbook(p, data_only=True), "two_stage")
                   for p in recalculated}
        result = {"input_digest": digest(cases), "oracle": subprocess.check_output(
            [executable, "--version"], text=True).strip(), "cases": outputs}
        save_json(root / "oracle.json", result)
        return result


def score_metrics(actual, expected):
    """Return strict coverage and per-cell scores; invalid cells never pass."""
    if not isinstance(actual, dict):
        actual = {}
    checks = []
    for metric, values in expected.items():
        got_values = actual.get(metric)
        correct_shape = isinstance(got_values, list) and len(got_values) == len(values)
        for i, wanted in enumerate(values):
            got = got_values[i] if correct_shape else None
            numeric = isinstance(got, (int, float)) and not isinstance(got, bool) and math.isfinite(got)
            checks.append({"metric": metric, "index": i,
                           "actual": got if numeric else None, "expected": wanted,
                           "ok": numeric and math.isclose(got, wanted, rel_tol=REL_TOL, abs_tol=ABS_TOL)})
    schema_ok = set(actual) == set(expected) and all(
        isinstance(actual[k], list) and len(actual[k]) == len(expected[k]) for k in expected)
    return {"ok": schema_ok and all(c["ok"] for c in checks), "schema_ok": schema_ok,
            "passed": sum(c["ok"] for c in checks), "total": len(checks), "checks": checks}


def grade(root):
    cases = load_json(root / "hidden-inputs.json")
    oracle = load_json(root / "oracle.json")
    if oracle["input_digest"] != digest(cases) or set(oracle["cases"]) != set(cases):
        raise ValueError("Oracle/input mismatch")
    result = {"input_digest": digest(cases), "oracle": oracle["oracle"],
              "tolerances": {"relative": REL_TOL, "absolute": ABS_TOL}, "arms": {}}
    for arm in ("burr", "python"):
        packet = root / arm
        submission = packet / "submission.py"
        if not submission.is_file():
            result["arms"][arm] = {"status": "missing_submission"}
            continue
        before = hashlib.sha256(submission.read_bytes()).hexdigest()
        results = {}
        with tempfile.TemporaryDirectory(prefix=f"pilot-grade-{arm}-") as tmp:
            tmp = Path(tmp)
            for case, inputs in cases.items():
                run_dir = tmp / case
                run_dir.mkdir()
                source, output = run_dir / "inputs.json", run_dir / "output.json"
                save_json(source, inputs)
                env = {**os.environ, "PYTHONPATH": str(packet / "vendor")}
                try:
                    completed = subprocess.run([sys.executable, "-S", str(submission),
                        "--inputs", str(source), "--output", str(output)], cwd=packet,
                        env=env, text=True, capture_output=True, timeout=45)
                    actual = load_json(output) if completed.returncode == 0 and output.is_file() else {}
                    scored = score_metrics(actual, oracle["cases"][case])
                    scored.update({"exit_code": completed.returncode,
                                   "stderr": completed.stderr[-2000:]})
                except (OSError, ValueError, subprocess.TimeoutExpired) as exc:
                    scored = score_metrics({}, oracle["cases"][case])
                    scored["error"] = str(exc)
                results[case] = scored
        if hashlib.sha256(submission.read_bytes()).hexdigest() != before:
            raise ValueError(f"Submission changed during grading: {arm}")
        delta_results = {}
        for case in cases:
            if case == "base":
                continue
            baseline = results["base"]["checks"]
            current = results[case]["checks"]
            wanted, actual = {}, {}
            for b, c in zip(baseline, current):
                key = b["metric"]
                wanted.setdefault(key, []).append(c["expected"] - b["expected"])
                actual.setdefault(key, []).append(
                    c["actual"] - b["actual"] if b["actual"] is not None and c["actual"] is not None else None)
            delta_results[case] = score_metrics(actual, wanted)
        result["arms"][arm] = {
            "status": "graded", "submission_sha256": before, "cases": results,
            "delta_cases": delta_results,
            "baseline_pass": results["base"]["ok"],
            "heldout_cases_passed": sum(r["ok"] for k, r in results.items() if k != "base"),
            "heldout_cases_total": len(cases) - 1,
            "cells_passed": sum(r["passed"] for r in results.values()),
            "cells_total": sum(r["total"] for r in results.values()),
            "delta_cells_passed": sum(r["passed"] for r in delta_results.values()),
            "delta_cells_total": sum(r["total"] for r in delta_results.values()),
        }
    save_json(root / "scores.json", result)
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("oracle", "score"))
    parser.add_argument("run_dir", type=Path)
    args = parser.parse_args()
    result = make_oracle(args.run_dir) if args.command == "oracle" else grade(args.run_dir)
    print(json.dumps({"command": args.command, "completed": True,
        "summary": {k: {m: v for m, v in arm.items() if m not in ("cases", "delta_cases")}
                    for k, arm in result.get("arms", {}).items()}}, indent=2))
