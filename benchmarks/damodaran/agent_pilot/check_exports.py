from pathlib import Path
import json,math,hashlib
from openpyxl import load_workbook
from benchmarks.damodaran.oracle import convert,libreoffice
from benchmarks.damodaran.agent_pilot.grade import score_metrics
root=Path('/private/tmp/burr-agent-pilot-20261007/frozen')
expected=json.loads((root/'oracle.json').read_text())['cases']['base']
report={}
for arm,relative in [('burr','output/baseline.xlsx'),('python','output/baseline-output.xlsx')]:
 source=root/arm/relative
 formulas=load_workbook(source)
 result=convert(libreoffice(),[source],root/'recalculated'/arm)[0]
 values=load_workbook(result,data_only=True)
 errors=[f'{sheet.title}!{c.coordinate}: {c.value}' for sheet in values for row in sheet for c in row if c.data_type=='e']
 formula_cells=[(sheet.title,c.coordinate) for sheet in formulas for row in sheet for c in row if c.data_type=='f']
 missing=[f'{s}!{c}' for s,c in formula_cells if values[s][c].value is None]
 actual={}
 if arm=='python':
  sheet=values['Outputs']
  for row in sheet.iter_rows(min_row=2):
   key=row[0].value
   actual[key]=[c.value for c in row[1:len(expected[key])+1]]
 else:
  sheet=values['Model']
  rows={row[0].value:row[0].row for row in sheet if row[0].value}
  for key,exp in expected.items():
   if key=='ev': actual[key]=[values['DCF']['B9'].value]
   elif key in ('equity_before_options','per_share','pv_terminal','terminal_value'):
    actual[key]=[sheet.cell(rows[key],6).value]
   else: actual[key]=[sheet.cell(rows[key],2+i).value for i in range(len(exp))]
 scored=score_metrics(actual,expected)
 report[arm]={'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'formula_cells':len(formula_cells),'errors':errors,'missing_formula_caches':missing,'source_metric_checks':scored}
 assert not errors and not missing and scored['ok'], report[arm]
 if arm=='burr':
  assert math.isclose(values['DCF']['B13'].value,expected['per_share'][0],rel_tol=1e-9,abs_tol=1e-7)
(root/'excel-checks.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps({k:{f:v for f,v in r.items() if f!='source_metric_checks'}|{'passed':r['source_metric_checks']['passed'],'total':r['source_metric_checks']['total']} for k,r in report.items()},indent=2))
