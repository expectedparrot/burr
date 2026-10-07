#!/usr/bin/env python3
"""Direct reconstruction of the fixed-switch, five-year FCFF model."""
import argparse
import json
import math
from pathlib import Path
import sys

# Resolve the supplied vendored library independently of the calling directory.
sys.path.insert(0, str(Path(__file__).resolve().parent / 'vendor'))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties

INPUTS = (
    'beta', 'capex0', 'cash', 'debt', 'debt_cost', 'depreciation0',
    'ebit0', 'high_years', 'options', 'prior_book_debt', 'prior_book_equity',
    'revenue0', 'risk_premium', 'riskfree', 'share_price', 'shares',
    'stable_growth', 'stable_roc', 'tax_rate', 'working_capital0',
    'working_capital_change0',
)


def model(x):
    """Return contract metrics in source-cell order, without cached answers."""
    if set(x) != set(INPUTS):
        raise ValueError('Input must contain precisely the required numeric fields')
    if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v)
           for v in x.values()):
        raise ValueError('All inputs must be finite numbers')
    if x['high_years'] != 5:
        raise ValueError('This contract fixes the high-growth horizon at five years')
    nopat0 = x['ebit0'] * (1 - x['tax_rate'])
    roc = nopat0 / (x['prior_book_debt'] + x['prior_book_equity'])
    reinvestment_rate = (x['capex0'] - x['depreciation0'] + x['working_capital_change0']) / nopat0
    growth = reinvestment_rate * roc
    ke = x['riskfree'] + x['beta'] * x['risk_premium']
    equity_weight = x['share_price'] * x['shares'] / (x['debt'] + x['share_price'] * x['shares'])
    debt_weight = 1 - equity_weight
    kd = x['debt_cost'] * (1 - x['tax_rate'])
    wacc = ke * equity_weight + kd * debt_weight
    wc_ratio = max(0, x['working_capital0'] / x['revenue0'])
    nopat, net_capex, delta_wc, fcff, pv_fcff = [], [], [], [], []
    for year in range(1, 6):
        nopat.append((nopat[-1] if nopat else nopat0) * (1 + growth))
        net_capex.append(x['capex0'] * (1 + growth)**year - x['depreciation0'] * (1 + growth)**year)
        delta_wc.append(x['revenue0'] * wc_ratio * ((1 + growth)**year - (1 + growth)**(year - 1)))
        fcff.append(nopat[-1] - net_capex[-1] - delta_wc[-1])
        pv_fcff.append(fcff[-1] / (1 + wacc)**year)
    stable_growth = x['stable_growth']
    terminal_nopat = nopat0 * (1 + growth)**x['high_years'] * (1 + stable_growth)
    terminal_wc = (x['revenue0'] * (1 + growth)**x['high_years'] * (1 + stable_growth)
                   - x['revenue0'] * (1 + growth)**x['high_years']) * wc_ratio
    terminal_net_capex = (stable_growth / x['stable_roc']) * terminal_nopat - terminal_wc
    terminal_fcff = terminal_nopat - terminal_net_capex - terminal_wc
    terminal_value = terminal_fcff / (wacc - stable_growth)
    pv_terminal = terminal_value / (1 + wacc)**x['high_years']
    ev = sum(pv_fcff) + pv_terminal
    equity = ev + x['cash'] - x['debt']
    result = {
        'after_tax_debt_cost': [kd], 'cost_of_equity': [ke],
        'debt_weight': [debt_weight], 'equity_before_options': [equity],
        'ev': [ev], 'fcff': fcff + [terminal_fcff], 'high_growth': [growth],
        'net_capex': net_capex + [terminal_net_capex],
        'nopat': nopat + [terminal_nopat],
        'per_share': [(equity - x['options']) / x['shares']],
        'pv_fcff': pv_fcff, 'pv_terminal': [pv_terminal],
        'terminal_value': [terminal_value], 'wacc': [wacc],
        'working_capital_change': delta_wc + [terminal_wc],
        'working_capital_ratio': [wc_ratio],
    }
    if not all(math.isfinite(v) for values in result.values() for v in values):
        raise ValueError('Inputs produce a non-finite valuation')
    return result


def export_excel(x, path):
    """Author a new live-formula workbook, without opening the source workbook."""
    book = openpyxl.Workbook()
    inputs = book.active
    inputs.title = 'Inputs'
    inputs.append(['Input', 'Value'])
    refs = {}
    for name in INPUTS:
        inputs.append([name, x[name]])
        refs[name] = f"Inputs!$B${inputs.max_row}"
    inputs['D1'] = 'Fixed source selections'
    for row, text in enumerate([
        'Public equity market weights; unchanged debt ratio in stable period',
        'CAPM cost of equity; beta and debt cost unchanged in stable period',
        'Fundamental growth, using current reinvestment rate and prior book capital',
        'Capex, depreciation and revenue grow with earnings',
        'Current working capital / revenue ratio, floored at zero',
        'Stable reinvestment = stable growth / stable return on capital',
        'Five explicit forecast years, followed by terminal year',
    ], 2):
        inputs.cell(row, 4, text)
    calc = book.create_sheet('Model')
    calc.append(['Derived assumption', 'Live formula'])
    expressions = [
        ('Current NOPAT', f"{refs['ebit0']}*(1-{refs['tax_rate']})"),
        ('Return on capital', f"B2/({refs['prior_book_debt']}+{refs['prior_book_equity']})"),
        ('Reinvestment rate', f"({refs['capex0']}-{refs['depreciation0']}+{refs['working_capital_change0']})/B2"),
        ('High growth', 'B3*B4'),
        ('Cost of equity', f"{refs['riskfree']}+{refs['beta']}*{refs['risk_premium']}"),
        ('Equity weight', f"{refs['share_price']}*{refs['shares']}/({refs['debt']}+{refs['share_price']}*{refs['shares']})"),
        ('Debt weight', '1-B7'),
        ('After-tax debt cost', f"{refs['debt_cost']}*(1-{refs['tax_rate']})"),
        ('WACC (both stages)', 'B6*B7+B9*B8'),
        ('Working capital ratio', f"MAX(0,{refs['working_capital0']}/{refs['revenue0']})"),
        ('Stable growth', refs['stable_growth']),
        ('Stable return on capital', refs['stable_roc']),
    ]
    for label, expression in expressions:
        calc.append([label, '=' + expression])
    calc['A15'] = 'Forecast metric'
    for col in range(2, 7):
        calc.cell(15, col, col - 1)
    calc['G15'] = 'Terminal year'
    for row, label in {16:'NOPAT', 17:'Net capital spending', 18:'Change in working capital', 19:'FCFF', 20:'Present value of FCFF'}.items():
        calc.cell(row, 1, label)
    for col in range(2, 7):
        letter = openpyxl.utils.get_column_letter(col)
        previous = openpyxl.utils.get_column_letter(col - 1)
        calc.cell(16, col, '=$B$2*(1+$B$5)' if col == 2 else f'={previous}16*(1+$B$5)')
        calc.cell(17, col, f"={refs['capex0']}*(1+$B$5)^{letter}$15-{refs['depreciation0']}*(1+$B$5)^{letter}$15")
        calc.cell(18, col, f"={refs['revenue0']}*$B$11*((1+$B$5)^{letter}$15-(1+$B$5)^({letter}$15-1))")
        calc.cell(19, col, f'={letter}16-{letter}17-{letter}18')
        calc.cell(20, col, f'={letter}19/(1+$B$10)^{letter}$15')
    calc['G16'] = f"=$B$2*(1+$B$5)^{refs['high_years']}*(1+$B$12)"
    calc['G18'] = f"=({refs['revenue0']}*(1+$B$5)^{refs['high_years']}*(1+$B$12)-{refs['revenue0']}*(1+$B$5)^{refs['high_years']})*$B$11"
    calc['G17'] = '=($B$12/$B$13)*G16-G18'
    calc['G19'] = '=G16-G17-G18'
    valuation = [
        ('Terminal value', 'G19/(B10-B12)'),
        ('PV terminal value', f"B23/(1+B10)^{refs['high_years']}"),
        ('PV explicit cash flows', 'SUM(B20:F20)'),
        ('Enterprise value', 'B24+B25'),
        ('Equity before options', f"B26+{refs['cash']}-{refs['debt']}"),
        ('Equity per share', f"(B27-{refs['options']})/{refs['shares']}"),
    ]
    for row, (label, expression) in enumerate(valuation, 23):
        calc.cell(row, 1, label)
        calc.cell(row, 2, '=' + expression)
    outputs = book.create_sheet('Outputs')
    outputs.append(['Contract metric', 'Value 1', 'Value 2', 'Value 3', 'Value 4', 'Value 5', 'Terminal / value 6'])
    scalar = {
        'after_tax_debt_cost':'B9', 'cost_of_equity':'B6', 'debt_weight':'B8',
        'equity_before_options':'B27', 'ev':'B26', 'high_growth':'B5',
        'per_share':'B28', 'pv_terminal':'B24', 'terminal_value':'B23',
        'wacc':'B10', 'working_capital_ratio':'B11',
    }
    series = {'fcff':19, 'net_capex':17, 'nopat':16, 'pv_fcff':20, 'working_capital_change':18}
    for metric in sorted(set(scalar) | set(series)):
        if metric in scalar:
            outputs.append([metric, '=Model!' + scalar[metric]])
        else:
            row = series[metric]
            columns = 'BCDEF' if metric == 'pv_fcff' else 'BCDEFG'
            outputs.append([metric] + [f'=Model!{col}{row}' for col in columns])
    for sheet in book:
        sheet.freeze_panes = 'B2'
        sheet.column_dimensions['A'].width = 31
        for col in 'BCDEFG':
            sheet.column_dimensions[col].width = 21
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.data_type == 'f' or isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.000000;[Red](#,##0.000000)'
                    if sheet.title == 'Inputs':
                        cell.font = Font(color='0000FF')
    inputs.column_dimensions['D'].width = 86
    book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inputs', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    x = json.loads(args.inputs.read_text())
    result = model(x)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook = args.output.parent / 'output' / (args.output.stem + '.xlsx')
    export_excel(x, workbook)
    args.output.write_text(json.dumps(result, indent=2, allow_nan=False) + '\n')


if __name__ == '__main__':
    main()
