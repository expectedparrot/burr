"""Development-only checks; reads source formulas, never used by submission.py."""
import ast
import json
import math
import operator
from pathlib import Path
import random
import re
import openpyxl
from submission import model, export_excel

# A deliberately small, independent evaluator for the formula grammar actually
# present in this source and in our newly generated workbook. IF is lazy and
# Excel's numeric-vs-text comparison ordering matters for unused forecast years.
TOKEN = re.compile(r'"[^"]*"|(?:[A-Za-z_][A-Za-z_0-9]*!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?')
class FormulaCheck:
    def __init__(self, book, overrides=None):
        self.book, self.overrides, self.cache = book, overrides or {}, {}
    def cell(self, sheet, address):
        address = address.replace('$', '')
        key = (sheet, address)
        if key in self.cache:
            return self.cache[key]
        v = self.overrides.get(key, self.book[sheet][address].value)
        if isinstance(v, str) and v.startswith('='):
            def replace(m):
                token=m.group()
                if token.startswith('"'):
                    return token
                s, addr = token.split('!') if '!' in token else (sheet, token)
                if ':' in addr:
                    a,b=addr.split(':')
                    return f'RANGE({s!r},{a!r},{b!r})'
                return f'CELL({s!r},{addr!r})'
            expr=TOKEN.sub(replace, v[1:]).replace('^','**')
            expr=re.sub(r'(?<![<>=!])=(?!=)', '==', expr)
            v=self.evaluate(ast.parse(expr, mode='eval').body)
        if v is None:
            v=0
        self.cache[key]=v
        return v
    def evaluate(self,n):
        if isinstance(n,ast.Constant): return n.value
        if isinstance(n,ast.BinOp):
            operations={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv,ast.Pow:operator.pow}
            return operations[type(n.op)](self.evaluate(n.left),self.evaluate(n.right))
        if isinstance(n,ast.UnaryOp):
            return -self.evaluate(n.operand) if isinstance(n.op,ast.USub) else self.evaluate(n.operand)
        if isinstance(n,ast.Compare):
            a,b=self.evaluate(n.left), self.evaluate(n.comparators[0])
            if isinstance(a,str) != isinstance(b,str):
                a,b=int(isinstance(a,str)),int(isinstance(b,str))
            operations={ast.Eq:operator.eq,ast.Lt:operator.lt,ast.Gt:operator.gt,ast.LtE:operator.le,ast.GtE:operator.ge,ast.NotEq:operator.ne}
            return operations[type(n.ops[0])](a,b)
        if isinstance(n,ast.Call):
            name=n.func.id
            if name=='IF':
                return self.evaluate(n.args[1] if self.evaluate(n.args[0]) else n.args[2])
            args=[self.evaluate(a) for a in n.args]
            if name=='CELL': return self.cell(*args)
            if name=='RANGE':
                s,a,b=args
                return [self.cell(s,c.coordinate) for row in self.book[s][f'{a}:{b}'] for c in row]
            if name=='SUM':
                return sum(v for a in args for v in (a if isinstance(a,list) else [a]) if isinstance(v,(int,float)))
            if name=='MAX': return max(args)
        raise ValueError(ast.dump(n))


def compare(actual,expected,where):
    assert actual.keys()==expected.keys(), where
    for metric,values in actual.items():
        assert len(values)==len(expected[metric]), (where,metric)
        for a,b in zip(values,expected[metric]):
            assert math.isfinite(a), (where,metric,a)
            assert math.isclose(a,b,rel_tol=2e-12,abs_tol=2e-9), (where,metric,a,b)


def main():
    source=openpyxl.load_workbook('source.xlsx',data_only=False)
    cache=openpyxl.load_workbook('source.xlsx',data_only=True)
    contract=json.loads(Path('contract.json').read_text())
    baseline=json.loads(Path('baseline.json').read_text())
    sheet=contract['sheet']
    cached={m:[cache[sheet][c].value for c in spec['source_cells']] for m,spec in contract['outputs'].items()}
    compare(model(baseline),cached,'baseline cached source')
    scenarios=[('baseline',baseline)]
    for key in baseline:
        if key=='high_years': continue
        for factor in (0.7,1.3):
            scenarios.append((f'{key} x {factor}',dict(baseline,**{key:baseline[key]*factor})))
    rng=random.Random(71943)
    for number in range(200):
        x={k:v if k=='high_years' else v*rng.uniform(0.5,1.5) for k,v in baseline.items()}
        scenarios.append((f'combined {number}',x))
    scenarios.extend([
        ('negative working capital',dict(baseline,working_capital0=-1000)),
        ('zero working capital',dict(baseline,working_capital0=0)),
        ('zero stable growth',dict(baseline,stable_growth=0)),
        ('negative stable growth',dict(baseline,stable_growth=-0.02)),
        ('negative high growth',dict(baseline,capex0=500,working_capital_change0=0)),
        ('zero debt',dict(baseline,debt=0)),
    ])
    for name,x in scenarios:
        overrides={(sheet,contract['input_fields'][k]):v for k,v in x.items()}
        evaluator=FormulaCheck(source,overrides)
        expected={m:[evaluator.cell(sheet,c) for c in spec['source_cells']] for m,spec in contract['outputs'].items()}
        compare(model(x),expected,name)
    excel_checks=[]
    for index in (0,49,len(scenarios)-1):
        name,x=scenarios[index]
        path=Path('output')/f'verification-{index}.xlsx'
        export_excel(x,path)
        book=openpyxl.load_workbook(path,data_only=False)
        evaluator=FormulaCheck(book)
        actual={}
        for row in range(2,book['Outputs'].max_row+1):
            metric=book['Outputs'].cell(row,1).value
            actual[metric]=[evaluator.cell('Outputs',book['Outputs'].cell(row,col).coordinate) for col in range(2,2+contract['outputs'][metric]['length'])]
        compare(actual,model(x),'new workbook '+name)
        excel_checks.append(name)
    report={
        'baseline_metrics_verified':len(cached),
        'source_formula_scenarios_verified':len(scenarios),
        'generated_live_workbooks_verified':excel_checks,
        'relative_tolerance':2e-12,'absolute_tolerance':2e-9,
        'excel_or_libreoffice_execution':False,
    }
    Path('verification.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))

if __name__=='__main__':main()
